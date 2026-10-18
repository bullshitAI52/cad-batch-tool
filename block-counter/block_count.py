"""Read-only block quantity reporting; DWG input needs ODA, DXF does not."""
import argparse
from collections import defaultdict
import csv
from datetime import datetime
import os
from pathlib import Path
import sys
import uuid
import math
from overlap import analyze, load_rules, SUMMARY_HEADERS, DETAIL_HEADERS

import ezdxf
from ezdxf.addons import odafc
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ['文件路径', '空间/布局', '有效图层', '原始块名', '类别', '直接数量', '嵌套数量']
STATUS_HEADERS = ['文件路径', '状态', '说明']
NOTE = '直接数量=布局中的块实例；嵌套数量=块内部引用乘以父级实例数。父块与子块不要相加作为设备总数。'


def configure_oda(path=None):
    if path:
        executable = Path(path).expanduser().resolve()
        if not executable.is_file():
            raise FileNotFoundError(f'ODA 程序不存在：{executable}')
        option = 'win_exec_path' if os.name == 'nt' else 'unix_exec_path'
        ezdxf.options.set('odafc-addon', option, str(executable))
    elif os.name == 'nt' and not odafc.is_installed():
        base = Path(os.environ.get('ProgramFiles', r'C:\Program Files')) / 'ODA'
        matches = sorted(base.glob('**/ODAFileConverter.exe'))
        if matches:
            ezdxf.options.set('odafc-addon', 'win_exec_path', str(matches[-1]))


def count_blocks(doc, source, scope='all', nested=True, max_depth=32, max_visits=100000):
    counts = defaultdict(lambda: [0, 0])
    notices = set()
    visits = 0
    truncated = False

    def walk(ref, layout, parent_layer, multiplier, chain):
        nonlocal visits, truncated
        if visits >= max_visits:
            truncated = True
            notices.add('达到安全遍历上限，数量不完整')
            return
        visits += 1
        name = ref.dxf.name
        identity = name.casefold()
        if identity in chain:
            notices.add(f'检测到循环块引用，已停止该分支：{name}')
            return
        if len(chain) > max_depth:
            notices.add('超过嵌套深度上限，深层数量未统计')
            return
        block = ref.block()
        category = '普通块'
        if block is None:
            category = '缺失定义'
            notices.add(f'找不到块定义：{name}')
        elif block.block.dxf.flags & 12:
            category = '外部参照'
            notices.add('外部参照只统计引用次数，不读取外部图纸内容')
        elif name.startswith('*'):
            category = '匿名/可能为动态块'
            notices.add('匿名/动态块保留原始名称，未合并为有效块名')
        layer = ref.dxf.get('layer', '0')
        if layer == '0' and chain:
            layer = parent_layer
        # Match ezdxf multi_insert(): zero-spacing rows/columns occupy one location.
        quantity = multiplier * max(1, int(ref.mcount))
        key = (str(source), layout, layer, name, category)
        counts[key][1 if chain else 0] += quantity
        if nested and block is not None and category != '外部参照':
            for child in block.query('INSERT'):
                if truncated:
                    break
                walk(child, layout, layer, quantity, chain + (identity,))

    layouts = [doc.modelspace()] if scope == 'model' else list(doc.layouts)
    for layout in layouts:
        for ref in layout.query('INSERT'):
            if truncated:
                break
            walk(ref, layout.name, '0', 1, ())
    rows = [list(key) + values for key, values in sorted(counts.items())]
    return rows, sorted(notices)


def load_drawing(path):
    if path.suffix.lower() == '.dxf':
        return ezdxf.readfile(path)
    if path.suffix.lower() == '.dwg':
        if not odafc.is_installed():
            raise RuntimeError('DWG 需要安装 ODA File Converter；也可先用 DXF 文件')
        return odafc.readfile(path, version='R2018', audit=False)
    raise ValueError('仅支持 .dwg 和 .dxf 文件')


def safe_csv(value):
    # Prevent drawing names/text being evaluated as spreadsheet formulas.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def write_sheet(ws, headers, rows, note, widths):
    ws.append([note])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.row_dimensions[1].height = 45
    ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical='center')
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            if cell.row > 2:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, int):
                    cell.number_format = '#,##0'
    for cell in ws[2]:
        cell.fill = PatternFill('solid', fgColor='234666')
        cell.font = Font(color='FFFFFF', bold=True)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'D3'
    ws.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{max(2,ws.max_row)}'


def export_reports(rows, statuses, output_dir, scope, nested, dedup=None, overlap_rows=None, tolerance=0.1):
    output_dir.mkdir(parents=True, exist_ok=True)
    base = output_dir / ('块数量统计_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid.uuid4().hex[:8])
    xlsx, csv_path, status_csv = [Path(str(base) + ext) for ext in ('.xlsx', '.csv', '_处理记录.csv')]
    wb = Workbook()
    detail = wb.active
    detail.title = '块数量统计'
    note = NOTE + f' 范围：{scope}；嵌套统计：{nested}；包括隐藏/关闭/冻结图层。'
    write_sheet(detail, HEADERS, rows, note, [46, 22, 24, 32, 24, 16, 16])
    write_sheet(wb.create_sheet('处理记录'), STATUS_HEADERS, statuses,
                '检查失败及警告：失败图纸不计入统计；警告可能表示统计不完整。', [55, 18, 95])
    extra = []
    if dedup is not None:
        summary = wb.create_sheet('分类去重统计', 0)
        write_sheet(summary, SUMMARY_HEADERS, dedup,
                    f'分类为规则匹配，需核对。只扣完全重复；疑似不扣。位置容差={tolerance} 图纸单位。直接与嵌套不可跨块名盲目相加。',
                    [46,20,24,24,30,16,16,16,16,18,18,18])
        write_sheet(wb.create_sheet('重复位置明细'), DETAIL_HEADERS, overlap_rows or [],
                    '位置为块基点的世界坐标；对象路径含父级句柄及阵列序号。疑似组按相邻关系连通，组内两端距离可能超过容差。',
                    [46,16,18,20,24,24,28,14,45,20,20,20,16,70])
        extra = [(Path(str(base)+'_去重统计.csv'),SUMMARY_HEADERS,dedup),
                 (Path(str(base)+'_重复明细.csv'),DETAIL_HEADERS,overlap_rows or [])]
    wb.save(xlsx)
    for path, headers, content in [(csv_path, HEADERS, rows), (status_csv, STATUS_HEADERS, statuses)] + extra:
        with path.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows([safe_csv(v) for v in row] for row in content)
    return (xlsx, csv_path, status_csv, *(item[0] for item in extra))


def run_batch(files, output_dir, scope='all', nested=True, deduplicate=True, tolerance=0.1,
              rules=(), progress=None, stop=None):
    """Shared CLI/UI worker. Callbacks must not directly access GUI widgets."""
    if not math.isfinite(tolerance) or tolerance <= 0:
        raise ValueError('位置容差必须大于零')
    paths = list(dict.fromkeys(Path(f).expanduser().resolve() for f in files))
    rows, statuses, dedup, overlap_rows = [], [], [], []
    completed = 0
    def emit(stage, **info):
        if progress:
            progress({'stage': stage, 'completed': completed, 'total': len(paths), **info})
    for path in paths:
        if stop is not None and stop.is_set():
            break
        emit('file_start', file=str(path))
        status = '成功'
        try:
            doc = load_drawing(path)
            emit('counting', file=str(path))
            items, notices = count_blocks(doc, path, scope, nested)
            rows.extend(items)
            if deduplicate:
                emit('deduplicating', file=str(path))
                try:
                    clean, positions, issues = analyze(doc, path, scope, nested, tolerance, rules)
                    dedup.extend(clean)
                    overlap_rows.extend(positions)
                    notices = sorted(set(notices + issues))
                except Exception as exc:
                    notices.append('去重分析失败，仅保留基础计数：' + str(exc))
            if notices:
                status = '警告'
                statuses.extend([[str(path), status, n] for n in notices])
            else:
                statuses.append([str(path), status, '已统计' if items else '未发现 INSERT 块引用（数量为零）'])
        except Exception as exc:
            status = '失败'
            statuses.append([str(path), status, str(exc)])
        completed += 1
        emit('file_done', file=str(path), status=status,
             messages=[r[2] for r in statuses if r[0] == str(path)])
    cancelled = completed < len(paths)
    reports = ()
    if completed:
        if cancelled:
            statuses.append(['本次任务', '已停止', f'仅处理 {completed}/{len(paths)} 份图纸，未处理图纸不计入结果'])
        emit('exporting')
        reports = export_reports(rows, statuses, Path(output_dir), scope, nested,
                                 dedup if deduplicate else None, overlap_rows, tolerance)
    return {'reports': reports, 'completed': completed, 'total': len(paths), 'cancelled': cancelled,
            'failed': len({r[0] for r in statuses if r[1] == '失败'}),
            'warnings': len({r[0] for r in statuses if r[1] == '警告'})}


def main():
    parser = argparse.ArgumentParser(description='只读统计 CAD 块数量，导出 Excel 和 CSV；无需 AutoCAD')
    parser.add_argument('files', nargs='*')
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--oda-path', default=os.environ.get('ODA_FILE_CONVERTER'))
    parser.add_argument('--scope', choices=['all', 'model'], default='all')
    parser.add_argument('--no-nested', action='store_true')
    parser.add_argument('--no-dedup', action='store_true', help='仅运行基础计数')
    parser.add_argument('--tolerance', type=float, default=0.1, help='疑似重叠位置容差，单位与图纸坐标相同')
    parser.add_argument('--rules', type=Path, default=Path(__file__).with_name('classification.json'))
    args = parser.parse_args()
    if not math.isfinite(args.tolerance) or args.tolerance<=0:
        parser.error('--tolerance 必须为大于零的有限数值')
    rules = load_rules(args.rules) if not args.no_dedup else []
    if not args.files:
        import tkinter as tk
        from tkinter.filedialog import askopenfilenames
        root = tk.Tk()
        root.withdraw()
        args.files = askopenfilenames(title='选择要统计的图纸（可多选）', filetypes=[('CAD', '*.dwg *.dxf')])
        root.destroy()
        if not args.files:
            return 0
    configure_oda(args.oda_path)
    output_dir = args.output_dir or Path(args.files[0]).resolve().parent / '块统计结果'
    def progress(event):
        if event['stage'] == 'file_start':
            print('统计：', event['file'], flush=True)
    result = run_batch(args.files, output_dir, args.scope, not args.no_nested,
                       not args.no_dedup, args.tolerance, rules, progress)
    for path in result['reports']:
        print('已生成：', path)
    return 1 if result['failed'] else 0


if __name__ == '__main__':
    sys.exit(main())
