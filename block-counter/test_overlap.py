import tempfile
from pathlib import Path
import unittest
import ezdxf
from openpyxl import load_workbook
from overlap import analyze
from block_count import export_reports


class OverlapTests(unittest.TestCase):
    def drawing(self):
        doc=ezdxf.new('R2018')
        doc.layers.new('灯具')
        doc.blocks.new('筒灯').add_circle((0,0),1)
        return doc

    def insert(self,doc,p=(0,0),**attrs):
        return doc.modelspace().add_blockref('筒灯',p,dxfattribs={'layer':'灯具',**attrs})

    def test_exact_nearby_rotation_and_far(self):
        doc=self.drawing()
        for p in [(0,0),(0,0),(0.05,0),(10,0)]:self.insert(doc,p)
        self.insert(doc,(0,0),rotation=90)
        rows,details,_=analyze(doc,'a.dxf',rules=[{'category':'灯具','block':'*灯*'}])
        row=rows[0]
        self.assertEqual(row[3],'灯具')
        self.assertEqual(row[5:],[5,0,1,0,4,0,1])
        self.assertEqual(sum(r[12]=='是' for r in details),1)

    def test_nested_world_transform_arrays_and_basepoint(self):
        doc=self.drawing()
        parent=doc.blocks.new('组件',base_point=(1,0))
        parent.add_blockref('筒灯',(2,0))
        for _ in range(2):
            doc.modelspace().add_blockref('组件',(10,0),dxfattribs={
                'layer':'灯具','rotation':90,'xscale':2,'yscale':2,'zscale':2,
                'column_count':2,'column_spacing':5})
        rows,details,_=analyze(doc,'a.dxf')
        lamp=next(r for r in rows if r[4]=='筒灯')
        self.assertEqual(lamp[5:11],[0,4,0,2,0,2])
        points=[r[9:12] for r in details if r[6]=='筒灯']
        self.assertTrue(any(abs(p[0]-10)<1e-9 and abs(p[1]-2)<1e-9 for p in points))

    def test_attributes_prevent_deduction(self):
        doc=self.drawing()
        for value in ('A','B'):
            self.insert(doc).add_attrib('编号',value,(0,0))
        rows,_,_=analyze(doc,'a.dxf')
        self.assertEqual(rows[0][7],0)
        self.assertEqual(rows[0][9],2)
        self.assertEqual(rows[0][-1],1)

    def test_layout_layer_and_z_separation(self):
        doc=self.drawing()
        self.insert(doc)
        self.insert(doc,(0,0,2))
        doc.layers.new('其他')
        self.insert(doc,layer='其他')
        doc.layouts.new('Sheet').add_blockref('筒灯',(0,0),dxfattribs={'layer':'灯具'})
        rows,details,_=analyze(doc,'a.dxf')
        self.assertEqual(sum(r[7]+r[8] for r in rows),0)
        self.assertFalse(details)

    def test_different_names_and_opaque_are_suspect_only(self):
        doc=self.drawing()
        self.insert(doc)
        doc.blocks.new('另一个灯').add_circle((0,0),1)
        doc.modelspace().add_blockref('另一个灯',(0,0),dxfattribs={'layer':'灯具'})
        doc.appids.new('CUSTOM')
        unknown=self.insert(doc)
        unknown.set_xdata('CUSTOM',[(1000,'metadata')])
        rows,details,issues=analyze(doc,'a.dxf')
        self.assertEqual(sum(r[7]+r[8] for r in rows),0)
        self.assertTrue(all(r[2]=='疑似重复' for r in details))
        self.assertTrue(issues)

    def test_grid_boundary(self):
        doc=self.drawing()
        self.insert(doc,(-0.01,0))
        self.insert(doc,(0.01,0))
        rows,_,_=analyze(doc,'a.dxf',tolerance=0.1)
        self.assertEqual(rows[0][-1],1)
        self.assertEqual(rows[0][7],0)

    def test_export_dedup_tabs(self):
        doc=self.drawing()
        self.insert(doc);self.insert(doc)
        rows,details,_=analyze(doc,'a.dxf')
        with tempfile.TemporaryDirectory() as td:
            paths=export_reports([],[],Path(td),'all',True,rows,details,0.1)
            self.assertEqual(len(paths),5)
            wb=load_workbook(paths[0])
            self.assertEqual(wb.sheetnames[0],'分类去重统计')
            self.assertEqual(wb['分类去重统计']['H3'].value,1)
            self.assertEqual(wb['分类去重统计']['J3'].value,1)
            self.assertEqual(wb['重复位置明细'].max_row,4)
            wb.close()


if __name__=='__main__':unittest.main()
