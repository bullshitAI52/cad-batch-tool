import csv
import hashlib
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import ezdxf
from openpyxl import load_workbook
from block_count import count_blocks, export_reports, load_drawing


class BlockTests(unittest.TestCase):
    def test_nested_arrays_layouts_and_inherited_layer(self):
        doc = ezdxf.new()
        doc.layers.new('设备')
        doc.blocks.new('螺栓').add_circle((0, 0), 1)
        parent = doc.blocks.new('组件')
        parent.add_blockref('螺栓', (0, 0), dxfattribs={'column_count': 2, 'column_spacing': 5})
        doc.modelspace().add_blockref('组件', (0, 0), dxfattribs={
            'layer': '设备', 'row_count': 2, 'column_count': 3, 'row_spacing': 10, 'column_spacing': 10})
        doc.modelspace().add_blockref('螺栓', (20, 0), dxfattribs={'layer': '设备'})
        doc.layouts.new('Sheet').add_blockref('螺栓', (0, 0))
        rows, warnings = count_blocks(doc, 'fixture.dxf')
        values = {(r[1], r[2], r[3]): r[-2:] for r in rows}
        self.assertEqual(values[('Model', '设备', '组件')], [6, 0])
        self.assertEqual(values[('Model', '设备', '螺栓')], [1, 12])
        self.assertEqual(values[('Sheet', '0', '螺栓')], [1, 0])
        self.assertFalse(warnings)
        direct, _ = count_blocks(doc, 'fixture.dxf', scope='model', nested=False)
        self.assertEqual(len(direct), 2)
        self.assertTrue(all(row[-1] == 0 for row in direct))

    def test_zero_spacing_and_repeated_definitions(self):
        doc = ezdxf.new()
        doc.blocks.new('A')
        doc.modelspace().add_blockref('A', (0, 0), dxfattribs={
            'row_count': 5, 'column_count': 3, 'column_spacing': 2})
        rows, _ = count_blocks(doc, 'file.dxf')
        self.assertEqual(rows[0][-2:], [3, 0])

    def test_cycle_missing_definition_and_xref(self):
        doc = ezdxf.new()
        a = doc.blocks.new('A')
        a.add_blockref('A', (0, 0))
        doc.modelspace().add_blockref('A', (0, 0))
        doc.modelspace().add_blockref('Missing', (0, 0))
        xref = doc.blocks.new('External')
        xref.block.dxf.flags = 4
        xref.add_blockref('A', (0, 0))
        doc.modelspace().add_blockref('External', (0, 0))
        rows, notices = count_blocks(doc, 'file.dxf')
        self.assertEqual(len(rows), 3)
        self.assertTrue(any('循环' in n for n in notices))
        self.assertTrue(any('找不到' in n for n in notices))
        self.assertTrue(any('外部' in n for n in notices))

    def test_source_and_export(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / '中文.dxf'
            doc = ezdxf.new()
            doc.blocks.new('设备')
            doc.modelspace().add_blockref('设备', (0, 0))
            doc.saveas(src)
            digest = hashlib.sha256(src.read_bytes()).digest()
            rows, _ = count_blocks(load_drawing(src), src)
            rows[0][3] = '=SUM(1,2)'  # Independently exercise exporter text safety.
            xlsx, csv_path, _ = export_reports(rows, [[str(src), '成功', '已统计']], Path(td), 'all', True)
            wb = load_workbook(xlsx)
            ws = wb['块数量统计']
            self.assertEqual(ws['D3'].value, '=SUM(1,2)')
            self.assertEqual(ws['D3'].data_type, 's')
            self.assertEqual(ws['F3'].value, 1)
            self.assertEqual(ws['F3'].data_type, 'n')
            self.assertEqual(ws.freeze_panes, 'D3')
            wb.close()
            with csv_path.open(encoding='utf-8-sig', newline='') as f:
                records = list(csv.reader(f))
            self.assertEqual(records[1][3], "'=SUM(1,2)")
            self.assertEqual(hashlib.sha256(src.read_bytes()).digest(), digest)

    def test_cli_partial_failure_and_duplicate_input(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / 'empty.dxf'
            ezdxf.new().saveas(src)
            out = Path(td) / 'out'
            result = subprocess.run([sys.executable, str(Path(__file__).with_name('block_count.py')),
                                     '--output-dir', str(out), str(src), str(src), str(Path(td)/'missing.dxf')],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 1, result.stderr)
            wb = load_workbook(next(out.glob('*.xlsx')))
            self.assertEqual(wb['处理记录'].max_row, 4)
            self.assertEqual(wb['处理记录']['B3'].value, '成功')
            self.assertEqual(wb['处理记录']['B4'].value, '失败')
            wb.close()


if __name__ == '__main__':
    unittest.main()
